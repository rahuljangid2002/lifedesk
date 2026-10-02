"""Builds the LifeDesk Demo Guide (.docx and .pdf) with screenshots taken in demo mode with the made-up sample data.

    python3 -m http.server 8765                 (in 4-Web-App)
    <venv>/python demo/build_guide.py           -> docs/LifeDesk - Demo Guide.docx / .pdf   (screenshots in docs/shots)
    <venv>/python demo/build_guide.py --no-shots   reuse the screenshots already taken
"""
import datetime, html, os, subprocess, sys, time

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), 'docs')
SHOTS = os.path.join(OUT, 'shots')
NAME = 'LifeDesk - Demo Guide'
URL = 'http://localhost:8765/'
DEMO = URL + '?demo=1'
TODAY = datetime.date.today().strftime('%-d %B %Y')
CHROME = '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'

INTRO = ('LifeDesk is a web app with one login and several everyday tools. The first tool is Money: accounts, entries, daily expenses, '
         'a monthly budget, people and loans, dashboards and reports. It works in any browser on a phone or a desktop, in any currency, '
         'and can be added to a phone\'s home screen.')
NOTE = ('Every name, amount and date in this guide is made-up sample data, loaded with the app\'s own "Load sample data" button in demo mode. '
        'No real account or real figures are shown.')

# (title, [screenshots], [points])   a screenshot name starting with "m_" is a phone screenshot
SECTIONS = [
    ('Signing in', ['01_login', '02_signup_code'], [
        'Sign in with an email and password, or with Google.',
        'Create an account: name, email and a password of at least 8 characters.',
        'A 6-digit code is sent to that email. The account is created only after the code is entered (10 minutes, 5 tries).',
        'Forgot password? sends a reset email. Every sign-in starts on the All tools screen.',
    ]),
    ('All tools', ['03_hub'], [
        'One login, several tools. Money is ready; Renewal reminders is shown as coming soon.',
        'The Money tile shows the net balance and this month\'s spending.',
        'On a wide screen the menu is on the left; on a phone there is a bar at the bottom.',
    ]),
    ('Money home', ['04_home'], [
        'This month\'s income, expense and net balance, with money in hand, card owed and money owed to people.',
        'Shortcuts to Add Entry, People & Loans, Budget, Daily, Dashboards and Reports.',
        'Accounts with their balances, budget progress, people, and the latest entries.',
        'Net balance = bank + wallets + cash − credit card owed − money owed to people.',
    ]),
    ('Add Entry', ['05_add', '06_edit'], [
        'Seven kinds of entry in one form: Expense, Income, Transfer, Lent, Got back, Borrowed, Repaid. A line under the title explains the chosen kind.',
        'Type the amount and a description: the category is suggested from the description.',
        'Account chips show each balance; a credit card shows what it owes.',
        'Recent entries are listed beside the form. The pencil opens an entry to change or delete it; balances, budget and people follow.',
    ]),
    ('Daily Expenses', ['07_daily'], [
        'Every entry date by date, with the total and the number of entries for each day.',
        'A month at a time, or any custom dates.',
        'Spending, Income or All entries; filter by category, account or text.',
        'Total, average per day and the highest day are in the header.',
    ]),
    ('Budget', ['08_budget', '09_budget_entries'], [
        'One row per category: what was spent this month, what is left, and the budget amount.',
        'Categories in use come first, largest budget on top; last month and the 3-month average are shown as hints.',
        '"See entries" opens the expenses behind a category.',
        'Copy last month or use the 3-month average. After a change, a bar at the bottom offers to save.',
    ]),
    ('Dashboards: Monthly', ['10_dash_monthly', '11_dash_monthly_charts'], [
        'Opening balance, closing balance and money in hand; income, expense and savings with the change against last month.',
        'Spend by category (largest first) and budget against actual, with the state written in words.',
        'Daily spend: point at a bar for the date and amount. Spend by account and the top 10 expenses.',
        'Every chart has "Show as table" for the same figures as a table.',
    ]),
    ('Dashboards: Yearly and Balances', ['12_dash_yearly', '13_dash_balances'], [
        'Yearly: income, expense, savings and savings rate; income against expense for each month; the net balance trend; top categories; income by source.',
        'The month a year starts in can be chosen (April for a financial year, January for a calendar year).',
        'Balances: money in each account, what each card owes, who owes you and whom you owe.',
    ]),
    ('Reports', ['14_report_category', '15_report_pivot'], [
        '15 reports as tables with totals: monthly summary, spend by category, category by month, daily spend, budget vs actual, top expenses, '
        'spend by account, income by source, year over year, lending activity, money to receive, owed to people, bad debts, account balances, all entries.',
        'Choose the period: this month, last month, last 3 months, this year, last year, all time, or custom dates.',
        'Download CSV opens in any spreadsheet.',
    ]),
    ('People & Loans', ['16_people'], [
        'Who owes you and whom you owe, with the amount still outstanding and a status: Active, Settled or Bad Debt.',
        'Open a person for quick actions (got money back, gave more, paid back, borrowed more), phone, notes and the full history.',
        'Bad debts are kept out of "to receive".',
    ]),
    ('My Accounts', ['17_accounts', '18_pay_card'], [
        'Cash and bank accounts, wallets and credit cards, grouped, each with its balance.',
        'A credit card shows how much of its limit is used and its bill due day.',
        'Pay bill: pay a card in full or in part from any account. The amount owed goes down; the net balance does not change.',
        'An account that is no longer used can be closed (history kept) or, if never used, deleted.',
    ]),
    ('Account, appearance and currency', ['19_account', '20_dark', '21_currency'], [
        'Appearance: System (follows the device), Light or Dark.',
        'Currency: any currency; amounts, number grouping and dates follow the region.',
        'Add your own categories with keywords for the suggestion.',
        'Download a backup file or restore one. Load or remove the sample data. Delete my account and data.',
    ]),
    ('On a phone', ['m_home', 'm_add', 'm_dash'], [
        'The same app with a bar at the bottom: Home, Daily, Add, Budget, Account.',
        'Open the address in the phone\'s browser and choose "Add to Home screen": it opens like any other app.',
        'The app\'s screens open without a connection; changes made offline are sent when the phone is back online.',
    ]),
]

FACTS = [
    ('Address', 'https://rahuljangid2002.github.io/lifedesk/  (try it without an account: add ?demo=1)'),
    ('Devices', 'Any modern browser on Android, iPhone, Windows or Mac'),
    ('Sign-in', 'Email + password with an emailed code at sign-up, or Google'),
    ('Privacy', 'Each user can reach only their own data'),
    ('Cost', 'Free to use; hosted on free services'),
    ('Not included yet', 'Buying on EMI, assets, recurring bills, receipts, mobile-number sign-in, renewal reminders'),
]


def take_shots():
    os.makedirs(SHOTS, exist_ok=True)
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        ctx = b.new_context(viewport={'width': 1440, 'height': 900}, device_scale_factor=2, color_scheme='light')
        pg = ctx.new_page()
        pg.on('dialog', lambda d: d.accept())
        shot = lambda name, page=None: (page or pg).screenshot(path=os.path.join(SHOTS, name + '.png')) and print('shot', name)
        go = lambda route, wait=0.6: (pg.evaluate("r => { location.hash = r; window.scrollTo(0, 0); }", route), time.sleep(wait))

        # sign-in screens (the real ones; the code request is answered here, so no email is sent and no account is made)
        ctx.route('**/macros/s/**', lambda r: r.fulfill(status=200, content_type='application/json', body='{"ok": true, "minutes": 10}'))
        pg.goto(URL); pg.wait_for_selector('.login-form', timeout=30000); time.sleep(0.5); shot('01_login')
        pg.get_by_role('button', name='Create an account').click()
        pg.get_by_label('Your name').fill('Aarav Mehta'); pg.get_by_label('Email').fill('aarav.mehta@example.com')
        pg.locator('[data-model="login.password"]').fill('a-long-password'); pg.get_by_role('button', name='Send verification code').click()
        pg.wait_for_selector('input.code'); pg.locator('input.code').fill('482913'); time.sleep(3.5); shot('02_signup_code')

        # demo mode, rupees, sample data
        pg.goto(DEMO); pg.wait_for_selector('#welcome-currency'); pg.locator('#welcome-currency').select_option('INR')
        pg.get_by_role('button', name='Continue').click(); pg.wait_for_selector('.tool')
        go('more'); pg.get_by_role('button', name='Load sample data').click(); pg.wait_for_selector('button:has-text("Remove sample data")'); time.sleep(4)
        go('hub'); shot('03_hub')
        go('home'); shot('04_home')
        go('add'); pg.get_by_label('Amount').fill('450'); pg.get_by_label('Description').fill('Coffee with team'); time.sleep(0.4); shot('05_add')
        pg.locator('.row .icon-btn').first.click(); time.sleep(0.6); shot('06_edit')
        pg.get_by_role('button', name='Cancel').first.click()
        go('daily'); pg.get_by_role('button', name='Previous month').click(); time.sleep(0.5); shot('07_daily')
        go('budget'); pg.get_by_role('button', name='Previous month').click(); time.sleep(0.5); shot('08_budget')
        pg.locator('.budget-row', has_text='Food & Dining').get_by_role('button', name='See entries').click(); time.sleep(0.4)
        pg.locator('.budget-row.open').scroll_into_view_if_needed(); pg.mouse.wheel(0, -160); time.sleep(0.4); shot('09_budget_entries')
        go('insights'); pg.get_by_role('button', name='Previous month').click(); time.sleep(0.6); shot('10_dash_monthly')
        pg.locator('.chart-card', has_text='Daily spend').scroll_into_view_if_needed(); pg.mouse.wheel(0, 120); time.sleep(0.4)
        pg.locator('.chart-card', has_text='Daily spend').locator('.hit').nth(1).hover(force=True); time.sleep(0.4); shot('11_dash_monthly_charts')
        pg.mouse.move(5, 5); pg.evaluate('window.scrollTo(0, 0)')
        pg.get_by_role('button', name='Yearly', exact=True).click(); time.sleep(0.6); shot('12_dash_yearly')
        pg.get_by_role('button', name='Balances', exact=True).click(); time.sleep(0.6); shot('13_dash_balances')
        go('reports'); pg.locator('[data-model="reports.id"]').select_option('spend-by-category'); pg.locator('[data-model="reports.preset"]').select_option('last-month'); time.sleep(0.5); shot('14_report_category')
        pg.locator('[data-model="reports.id"]').select_option('category-by-month'); time.sleep(0.5); shot('15_report_pivot')
        go('people'); pg.locator('.p-head', has_text='Alex').click(); time.sleep(0.5); shot('16_people')
        go('accounts'); shot('17_accounts')
        pg.get_by_role('button', name='Pay bill').click(); pg.get_by_role('button', name='Partial').click()
        pg.locator('.a-panel input[type=number]').fill('5000'); pg.locator('.a-panel').scroll_into_view_if_needed(); time.sleep(0.4); shot('18_pay_card')
        pg.get_by_role('button', name='Cancel').click()
        go('more'); shot('19_account')
        pg.get_by_role('button', name='Dark', exact=True).click(); go('home'); shot('20_dark')
        go('more'); pg.get_by_role('button', name='Light', exact=True).click()
        (pg.locator('[data-currency]').select_option('USD'), pg.get_by_role('button', name='Keep the numbers').click()); time.sleep(0.5); go('accounts'); shot('21_currency')
        go('more'); (pg.locator('[data-currency]').select_option('INR'), pg.get_by_role('button', name='Keep the numbers').click()); time.sleep(0.5)
        pg.get_by_role('button', name='System', exact=True).click()
        state = ctx.storage_state()
        ctx.close()

        m = b.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=3, storage_state=state, color_scheme='light')
        mp = m.new_page()
        for route, name in (('home', 'm_home'), ('add', 'm_add'), ('insights', 'm_dash')):
            mp.goto(DEMO + '#' + route); time.sleep(0.9); shot(name, mp)
        b.close()


def build_docx(path):
    navy, blue, grey = RGBColor(0x0B, 0x13, 0x24), RGBColor(0x25, 0x63, 0xEB), RGBColor(0x66, 0x70, 0x85)
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.6)
    d.styles['Normal'].font.name = 'Helvetica'; d.styles['Normal'].font.size = Pt(10.5)
    t = d.add_paragraph(); r = t.add_run('LifeDesk'); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = navy
    t = d.add_paragraph(); r = t.add_run('Demo Guide: one login, your everyday tools. Money first.'); r.font.size = Pt(15); r.font.color.rgb = blue
    t = d.add_paragraph(); r = t.add_run(f'Rahul Jangid · {TODAY}'); r.font.color.rgb = grey
    d.add_paragraph(INTRO)
    p = d.add_paragraph(); r = p.add_run(NOTE); r.italic = True; r.font.color.rgb = grey
    d.add_picture(os.path.join(SHOTS, '04_home.png'), width=Cm(17.4))
    for i, (title, shots, points) in enumerate(SECTIONS, start=1):
        d.add_page_break()
        h = d.add_heading(f'{i}. {title}', level=1)
        for run in h.runs:
            run.font.color.rgb = navy
        for point in points:
            d.add_paragraph(point, style='List Bullet')
        phone = [s for s in shots if s.startswith('m_')]
        if phone:
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for s in phone:
                p.add_run().add_picture(os.path.join(SHOTS, s + '.png'), width=Cm(5.3)); p.add_run('  ')
        for s in [s for s in shots if not s.startswith('m_')]:
            d.add_picture(os.path.join(SHOTS, s + '.png'), width=Cm(17.4 if len(shots) == 1 else 14.6))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()
    h = d.add_heading(f'{len(SECTIONS) + 1}. At a glance', level=1)
    for run in h.runs:
        run.font.color.rgb = navy
    table = d.add_table(rows=0, cols=2); table.style = 'Table Grid'
    for a, b in FACTS:
        row = table.add_row().cells; row[0].text = a; row[1].text = b
        row[0].paragraphs[0].runs[0].bold = True
    d.add_paragraph(); p = d.add_paragraph(); r = p.add_run(NOTE); r.italic = True; r.font.color.rgb = grey
    d.save(path)


def build_pdf(path):
    e = html.escape
    parts = [f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 15mm 14mm; }}
body {{ font-family: -apple-system, 'Segoe UI', Helvetica, Arial, sans-serif; color: #101828; font-size: 10.5pt; line-height: 1.45; }}
.cover {{ background: linear-gradient(135deg,#0B1324,#1B2A4A 60%,#2563EB); color: #fff; border-radius: 14px; padding: 30px 34px; margin-bottom: 16px; }}
.cover h1 {{ font-size: 34pt; margin: 0 0 6px; }} .cover p {{ margin: 4px 0; color: #C7D7FE; font-size: 13pt; }}
h2 {{ color: #0B1324; border-bottom: 2px solid #2563EB; padding-bottom: 4px; margin: 0 0 8px; }}
.sec {{ page-break-before: always; }}
.note {{ color: #667085; font-style: italic; }}
ul {{ margin: 6px 0 10px; padding-left: 18px; }} li {{ margin: 3px 0; }}
img {{ display: block; width: 100%; border: 1px solid #D8DDE6; border-radius: 8px; margin: 8px auto; }}
img.pair {{ width: 82%; }}
img.trio {{ width: 54%; margin: 5px auto; }}
.phones {{ display: flex; gap: 4%; justify-content: center; }} .phones img {{ width: 29%; margin: 8px 0; }}
table {{ border-collapse: collapse; width: 100%; }} td {{ border: 1px solid #D8DDE6; padding: 7px 9px; vertical-align: top; }} td:first-child {{ font-weight: 700; width: 24%; background: #EAF1FF; }}
</style></head><body>
<div class="cover"><h1>LifeDesk</h1><p>Demo Guide: one login, your everyday tools. Money first.</p><p>Rahul Jangid · {e(TODAY)}</p></div>
<p>{e(INTRO)}</p><p class="note">{e(NOTE)}</p><img src="file://{SHOTS}/04_home.png">"""]
    for i, (title, shots, points) in enumerate(SECTIONS, start=1):
        parts.append(f'<div class="sec"><h2>{i}. {e(title)}</h2><ul>' + ''.join(f'<li>{e(p)}</li>' for p in points) + '</ul>')
        phone = [s for s in shots if s.startswith('m_')]
        if phone:
            parts.append('<div class="phones">' + ''.join(f'<img src="file://{SHOTS}/{s}.png">' for s in phone) + '</div>')
        desk = [s for s in shots if not s.startswith('m_')]
        for s in desk:
            parts.append(f'<img class="{"trio" if len(desk) > 2 else "pair" if len(desk) > 1 else ""}" src="file://{SHOTS}/{s}.png">')
        parts.append('</div>')
    parts.append(f'<div class="sec"><h2>{len(SECTIONS) + 1}. At a glance</h2><table>' + ''.join(f'<tr><td>{e(a)}</td><td>{e(b)}</td></tr>' for a, b in FACTS)
                 + f'</table><p class="note">{e(NOTE)}</p></div></body></html>')
    src = os.path.join(OUT, 'src', 'demo_guide.html')
    os.makedirs(os.path.dirname(src), exist_ok=True)
    with open(src, 'w') as f:
        f.write(''.join(parts))
    subprocess.run([CHROME, '--headless', '--disable-gpu', '--no-pdf-header-footer', f'--print-to-pdf={path}', 'file://' + src], check=True, capture_output=True)


if __name__ == '__main__':
    if '--no-shots' not in sys.argv:
        take_shots()
    build_docx(os.path.join(OUT, NAME + '.docx'))
    build_pdf(os.path.join(OUT, NAME + '.pdf'))
    print('written', os.path.join(OUT, NAME + '.docx'), 'and .pdf')
